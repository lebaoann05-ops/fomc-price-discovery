"""Build the editable Word version of the Group 12 report from its Markdown source."""

from pathlib import Path
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt


SOURCE = Path(__file__).with_name("Group12_FOMC_Report.md")
TARGET = SOURCE.with_suffix(".docx")


def plain(text: str) -> str:
    text = re.sub(r"\[([^]]+)\]\(([^)]+)\)", r"\1 (\2)", text)
    text = text.replace("**", "").replace("`", "").replace("*", "")
    return text


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.8)
section.bottom_margin = Inches(0.75)
section.left_margin = Inches(0.9)
section.right_margin = Inches(0.9)

normal = doc.styles["Normal"]
normal.font.name = "Aptos"
normal.font.size = Pt(10.5)
normal.paragraph_format.space_after = Pt(7)
for name, size in [("Title", 17), ("Heading 1", 12.5), ("Heading 2", 11)]:
    doc.styles[name].font.name = "Aptos Display" if name == "Title" else "Aptos"
    doc.styles[name].font.size = Pt(size)

lines = SOURCE.read_text(encoding="utf-8").splitlines()
paragraph_lines: list[str] = []
table_lines: list[str] = []


def flush() -> None:
    if paragraph_lines:
        doc.add_paragraph(plain(" ".join(paragraph_lines)))
        paragraph_lines.clear()


def flush_table() -> None:
    if not table_lines:
        return
    rows = []
    for line in table_lines:
        cells = [plain(cell.strip()) for cell in line.strip().strip("|").split("|")]
        if all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells):
            continue
        rows.append(cells)
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    table.autofit = True
    for row_index, cells in enumerate(rows):
        for col_index, cell_text in enumerate(cells):
            cell = table.cell(row_index, col_index)
            cell.text = cell_text
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(2)
                for run in paragraph.runs:
                    run.font.size = Pt(8)
                    if row_index == 0:
                        run.bold = True
    doc.add_paragraph()
    table_lines.clear()


for line in lines:
    stripped = line.strip()
    if stripped.startswith("|"):
        flush()
        table_lines.append(stripped)
        continue
    flush_table()
    if not stripped:
        flush()
        continue
    if stripped.startswith("# "):
        flush()
        doc.add_paragraph(plain(stripped[2:]), style="Title")
    elif stripped.startswith("## "):
        flush()
        doc.add_heading(plain(stripped[3:]), level=1)
    elif stripped.startswith("### "):
        flush()
        doc.add_heading(plain(stripped[4:]), level=2)
    elif stripped.startswith("**Table "):
        flush()
        caption = doc.add_paragraph()
        caption.paragraph_format.keep_with_next = True
        caption.add_run(plain(stripped)).bold = True
    elif stripped.startswith("- "):
        flush()
        doc.add_paragraph(plain(stripped[2:]), style="List Bullet")
    elif stripped.endswith("  "):
        flush()
        doc.add_paragraph(plain(stripped))
    else:
        paragraph_lines.append(stripped)
flush()
flush_table()

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer.add_run("TCH442 · Group 12 · Report draft")
doc.save(TARGET)
print(TARGET)
